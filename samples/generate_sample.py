"""Generates samples/sample-questionnaire.xlsx (columns: Section | Question | Type)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

QUESTIONS = [
    ("Company Information", "Legal company name, headquarters location and years in business", "text"),
    ("Company Information", "Describe the product or service being evaluated and how it will be used at Contoso Health", "longtext"),
    ("Company Information", "Will the solution create, receive, store or transmit PHI?", "yesno"),
    ("Security Governance", "Do you maintain a documented information security program?", "yesno"),
    ("Security Governance", "Which security certifications or attestations do you hold (SOC 2, ISO 27001, HITRUST)?", "text"),
    ("Security Governance", "Do you perform third-party penetration tests at least annually?", "yesno"),
    ("Security Governance", "Describe your vulnerability management and patching process", "longtext"),
    ("Access Control", "Does the solution support multi-factor authentication?", "yesno"),
    ("Access Control", "Describe how user access is provisioned, reviewed and revoked", "longtext"),
    ("Access Control", "Is role-based access control enforced for administrative functions?", "yesno"),
    ("Data Protection", "Is data encrypted at rest (state algorithm and key length)?", "text"),
    ("Data Protection", "Is data encrypted in transit (state protocol versions)?", "text"),
    ("Data Protection", "Describe data retention, deletion and return processes at contract end", "longtext"),
    ("Data Protection", "Where is customer data stored and processed (countries/regions)?", "text"),
    ("Incident Response", "Do you have a documented incident response plan that is tested annually?", "yesno"),
    ("Incident Response", "What is your commitment for notifying customers of a security breach?", "text"),
    ("Business Continuity", "Do you have disaster recovery plans with defined RTO and RPO?", "yesno"),
    ("Business Continuity", "Describe backup frequency, location and restore testing", "longtext"),
    ("Privacy & Compliance", "Are you willing to sign a HIPAA Business Associate Agreement?", "yesno"),
    ("Privacy & Compliance", "Do you use subcontractors or fourth parties that access Contoso data? List them", "longtext"),
    ("Privacy & Compliance", "Is employee security awareness training performed at least annually?", "yesno"),
]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Questionnaire"
    ws.append(["Section", "Question", "Type"])
    for row in QUESTIONS:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in enumerate((24, 95, 12), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    out = Path(__file__).parent / "sample-questionnaire.xlsx"
    wb.save(out)
    print(f"Wrote {out} ({len(QUESTIONS)} questions)")


if __name__ == "__main__":
    main()
